from copy import deepcopy
from pathlib import Path
import shutil
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table


def first_run_rpr(paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def set_paragraph_text(paragraph, text):
    rpr = first_run_rpr(paragraph)
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, deepcopy(rpr))


def paragraph_by_text(doc, text):
    matches = [p for p in doc.paragraphs if p.text == text]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph {text!r}, found {len(matches)}")
    return matches[0]


def replace_exact(doc, old, new):
    set_paragraph_text(paragraph_by_text(doc, old), new)


def insert_after(anchor, text):
    new_p = OxmlElement("w:p")
    if anchor._p.pPr is not None:
        new_p.append(deepcopy(anchor._p.pPr))
    anchor._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    paragraph = Paragraph(new_p, anchor._parent)
    run = paragraph.add_run(text)
    rpr = first_run_rpr(anchor)
    if rpr is not None:
        run._r.insert(0, deepcopy(rpr))
    return paragraph


def insert_many_after(anchor, texts):
    current = anchor
    for text in texts:
        current = insert_after(current, text)
    return current


def clone_format(target, exemplar):
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if exemplar._p.pPr is not None:
        target._p.insert(0, deepcopy(exemplar._p.pPr))
    rpr = first_run_rpr(exemplar)
    for run in target.runs:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        if rpr is not None:
            run._r.insert(0, deepcopy(rpr))


def set_cell_text(cell, text):
    base_p = cell.paragraphs[0]
    rpr = first_run_rpr(base_p)
    ppr = deepcopy(base_p._p.pPr) if base_p._p.pPr is not None else None
    for p in list(cell.paragraphs)[1:]:
        cell._tc.remove(p._p)
    for run in list(base_p.runs):
        base_p._p.remove(run._r)
    if base_p._p.pPr is not None:
        base_p._p.remove(base_p._p.pPr)
    if ppr is not None:
        base_p._p.insert(0, ppr)
    run = base_p.add_run(text)
    if rpr is not None:
        run._r.insert(0, deepcopy(rpr))


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    header = trpr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        trpr.append(header)
    header.set(qn("w:val"), "true")


def prevent_fixed_row_height(row):
    trpr = row._tr.get_or_add_trPr()
    for height in list(trpr.findall(qn("w:trHeight"))):
        trpr.remove(height)


def insert_improvement_section(doc):
    target = paragraph_by_text(doc, "PRD B — Social Management")
    h1_exemplar = paragraph_by_text(doc, "13. 风险")
    normal_exemplar = paragraph_by_text(doc, "不是“生成更多任务”，而是“完成一个复活动作”。")

    heading = target.insert_paragraph_before("14. 功能改进清单与全部建议落地")
    clone_format(heading, h1_exemplar)
    intro = target.insert_paragraph_before(
        "本节把执行摘要、跨项目一致反馈、TryRevive 专项反馈和开放问题中的建议统一映射到产品功能、优先级与验证方式；未达到行为门槛前，不以新增功能数量作为进展。"
    )
    clone_format(intro, normal_exemplar)

    source_table = doc.tables[2]
    table_xml = deepcopy(source_table._tbl)
    target._p.addprevious(table_xml)
    table = Table(table_xml, doc)

    rows = [
        ("定位边界", "只服务“停滞至少 7 天、用户仍想继续”的真实项目；不以 AI 拆任务、通用待办或长计划为卖点。", "127；竞品定位；执行摘要", "P0；首屏与招募文案一致"),
        ("首发人群", "首轮主样本只选非技术学生/学生创业者；独立创作者与执行功能困难用户仅作后续或独立对照，不混入首发结论。", "127、130；开放问题", "P0；招募样本标签 100% 完整"),
        ("项目准入与继续判断", "在 intake 中确认停滞时长、继续理由、最后进展和可用时间，并允许“继续、缩小、暂不继续”三种正确结果。", "130；问题定义", "P0；2 分钟内完成"),
        ("上下文恢复", "优先利用用户主动选择的项目历史、最后产出和卡点；达到门槛后只接入一种文档来源，不同时做多平台导入。", "121；共同产品引擎；路线图", "P0 手工/P1 单一导入"),
        ("首个复活动作", "动作必须包含动词、对象、5-10 分钟上限和可验证完成标准，目标是先做到 60 分并产生可见结果。", "125、137；验收标准", "P0；过大动作率持续监控"),
        ("建议理由与用户控制", "显示建议依据和为什么是当前一步；用户可一键改小、换一步、拒绝，并用非技术语言描述结果和工具入口。", "关键用户故事；131", "P0；拒绝原因可记录"),
        ("Start Now", "单任务页只显示当前动作、计时器和真实文件/工具入口，不展示完整 backlog，避免继续规划。", "125、127；价值主张", "P0；计划停留率纳入护栏"),
        ("短冲刺模式", "完成首个微动作后，可进入最多 4 个 45 分钟阶段；每阶段只有一个交付物，默认只显示当前阶段。", "141", "P0 concierge 验证；产品化待门槛"),
        ("完成证据", "支持文本、链接、截图、文件或轻量自报；证据用于恢复上下文而非打卡审判，允许跳过并记录原因。", "开放问题；MVP 闭环", "P0；监控证据负担率"),
        ("D7 连续推进", "24 小时或 7 天后从最新证据继续，核心看第二次真实推进，而不是生成任务数或打开 App 次数。", "执行摘要；竞品定位", "P0；D7 二次推进率 ≥30%"),
        ("承诺机制", "在 concierge 中比较“AI 给第一步”和“AI + 真人/同伴承诺”；不采用羞耻、惩罚、贷款或财务押注。", "137；开放问题", "实验；验证后再产品化"),
        ("范围控制", "先跑通 intake—诊断—单动作—证据—回访；原生 App、协作、自动执行、项目管理和复杂积分全部后置。", "132；MVP 优先级", "P0；每次迭代只改核心闭环"),
        ("外部行为验证", "发起人自用满意只形成假设；必须记录外部用户是否当场开始、完成、拒绝或改小动作，以及失败原因。", "125、131、135", "P0；至少 12 次完整 concierge"),
        ("人工与自动建议对照", "比较人工建议与自动建议的完成率、修改量和上下文错配率，以外部结果证明建议质量。", "131；两周验证方案", "第 8-14 天；保留同一口径"),
        ("交付节奏", "3-7 天内交付可用 demo；一周内看不到真实行为就缩小场景。Hackathon 可作为期限和首批招募渠道。", "132、138", "执行要求；不以功能数为成果"),
        ("公开构建与获客", "Build in public 只用于招募和分享学习，不把关注量当作需求、留存或付费证明；同时记录渠道和招募成本。", "132；获客质疑", "实验看板单列渠道质量"),
        ("平台顺序与复制风险", "先做人工服务和轻量 Web；达到行为门槛后再评估 iOS。被复制风险不作为延期验证的理由。", "125、132", "P2 iOS；门槛前不原生开发"),
        ("数据与外部操作", "只读取用户选择的上下文，资料可查看、导出和删除；任何文件修改、消息发送或代码提交都必须显式确认。", "共同产品引擎；验收标准", "P0 护栏；未确认写入必须为 0"),
        ("商业化验证", "单独测试一次性“项目复活包”与订阅；不能用一般 iOS/安卓付费习惯推断本产品付费意愿。", "125；不能直接下结论", "30-90 天；记录真实支付/承诺"),
        ("命名与发布", "公开发布前检查 TryRevive 的域名、App Store、商标及健康/影像产品混淆风险。", "风险章节；90 天路线图", "发布前阻断项"),
        ("继续/调整/停止规则", "招募 20 名同一主样本，至少 12 人完成；会话成功率 ≥50%、完成者 D7 ≥30%。未达到先改场景/服务，仍不达标则停止。", "指标框架；停止标准", "第 14 天作明确决策"),
    ]

    needed = len(rows) + 1
    while len(table.rows) < needed:
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    while len(table.rows) > needed:
        table._tbl.remove(table.rows[-1]._tr)

    headers = ("改进项", "落地方式", "依据", "优先级 / 验收")
    for c, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[c], text)
    for r_idx, row_data in enumerate(rows, start=1):
        for c, text in enumerate(row_data):
            set_cell_text(table.rows[r_idx].cells[c], text)
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        prevent_fixed_row_height(row)


def main():
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)
    doc = Document(output)

    replacements = {
        "TryRevive 不能只做 AI 拆任务。 Goblin Tools 和 Tiimo 已经覆盖拆解、估时、排序、计时和下一步。真正值得测试的差异是：针对已经停滞至少 7 天的真实项目，保留其历史和卡点，在 10 分钟内让用户完成一个可验证的“复活动作”。":
            "TryRevive 不能只做 AI 拆任务。Goblin Tools 和 Tiimo 已经覆盖拆解、估时、排序、计时和下一步。真正值得测试的差异是：针对已经停滞至少 7 天的真实项目，保留其历史和卡点，在 10 分钟内让用户完成一个可验证的“复活动作”。首版还必须补齐继续/停止判断、建议理由、轻量证据和可选短冲刺，而不是扩展 backlog。",
        "帮助有停滞副项目的学生、独立创作者和非技术 Maker，在 10 分钟内找回上下文并完成一个真实的复活动作。":
            "帮助已有可展示副项目、但停滞至少 7 天的非技术学生/学生创业者，在 10 分钟内找回上下文并完成一个真实的复活动作。",
        "首发用户： 18-30 岁、手上有副项目、能使用 AI 工具但经常停在 0 到 1 的学生/独立创作者/非技术 Maker。":
            "首轮验证主用户：18-30 岁、有至少一个停滞 7 天以上副项目、能使用 AI 工具但经常卡在重新启动的非技术学生/学生创业者。独立创作者与执行功能困难用户作为后续或独立对照样本，不与首发样本混合下结论。",
        "选择要复活的项目。 输入项目名、原始目标、最后进展、停滞时间和当前可用时间。":
            "项目准入与继续判断。输入项目名、原始目标、最后进展、停滞时间和当前可用时间；确认项目是否仍值得继续，并允许“继续、缩小、暂不继续”三种结果。",
        "诊断卡点。 只问最多 3 个问题：为什么值得继续、最后一步是什么、现在最大的阻力是什么。":
            "诊断卡点。最多问 3 个问题：为什么值得继续、最后一步是什么、现在最大的阻力是什么；将阻力归为上下文缺失、动作过大、工具入口丢失或承诺不足，便于选择干预方式。",
        "生成 Revival Brief。 输出当前状态、一个 5-10 分钟动作、完成标准和后续两步。":
            "生成 Revival Brief。输出当前状态、一个 5-10 分钟动作、可验证完成标准、建议依据和后续两步；默认只展开当前动作，目标是先做到 60 分并产生可见结果。",
        "Start Now。 单任务界面、计时器、打开真实工具/文件的链接；不显示完整 backlog。":
            "Start Now。单任务界面、计时器、打开真实工具/文件的链接；不显示完整 backlog。完成首个微动作后，可选择进入最多 4 个 45 分钟短冲刺，每阶段只显示一个交付物。",
        "提交完成证据。 文本、链接、截图或文件；用户可以标记动作不合适并说明原因。":
            "提交完成证据。支持文本、链接、截图、文件或轻量自报；用户可跳过证据、标记动作不合适并说明原因，证据用于下次恢复上下文而不是制造打卡压力。",
        "导入 Markdown、Notion 导出或 GitHub README":
            "单一来源导入（首选 Markdown/README；Notion 导出与 GitHub README 只择一验证）",
        "分享式 build log":
            "分享式 build log（仅用于招募与分享学习，不作为需求/留存证明）",
        "完整项目管理、甘特图、复杂积分体系":
            "完整项目管理、甘特图、复杂积分体系，以及羞耻、惩罚或财务押注式督促",
        "用户把它当作另一个待办工具，缺少回访动机。":
            "定位稀释：用户把它当作另一个待办工具，缺少回访动机。缓解：只接收停滞至少 7 天的真实项目，默认只展示一个复活动作。",
        "建议动作看似合理但与真实上下文不符。":
            "上下文错配：建议动作看似合理但与真实现状不符。缓解：显示建议依据，允许改小、换一步或拒绝，并记录原因。",
        "游戏化奖励替代真实产出。":
            "伪进展：游戏化奖励替代真实产出。缓解：奖励只绑定真实完成证据与 D7 第二次推进，不为规划、打开 App 或生成任务给分。",
        "名称 TryRevive 与其他健康/影像产品混淆；公开发布前需检查域名、App Store 和商标。":
            "名称混淆：TryRevive 可能与健康/影像产品混淆。缓解：公开发布前检查域名、App Store、商标和平台政策。",
        "第一轮临时门槛（不是市场基准）： 招募 20 人；至少 12 人完成 concierge 会话；复活会话成功率 ≥50%；完成者 D7 二次推进率 ≥30%；若达不到，先改场景/服务，不开发原生 App。":
            "第一轮临时门槛（不是市场基准）：招募 20 名同一首发人群；至少 12 人完成 concierge 会话；复活会话成功率 ≥50%；完成者 D7 二次推进率 ≥30%；若达不到，先改场景/服务并复测，仍不达标则停止，不开发原生 App。",
        "找 20 个至少有一个停滞副项目的人。":
            "从同一首发人群招募 20 名至少有一个停滞 7 天以上副项目的人；其他人群单独标记，只作探索。",
        "现场观察用户是否立即行动；记录卡住的位置和拒绝原因。":
            "现场观察用户是否立即行动；记录卡住的位置、拒绝/改小原因、证据负担，以及是否需要真人/同伴承诺。",
        "比较人工建议与自动建议的完成率。":
            "比较人工建议与自动建议的完成率、修改量和上下文错配率；同时小样本比较纯 AI 与可选真人/同伴承诺、必填证据与可选证据。",
        "只实现 intake、单动作页、完成证据和下一次提醒。":
            "只实现 intake、继续/停止判断、单动作页、建议依据、轻量证据和下一次提醒；短冲刺先用 concierge 方式验证。",
        "开放问题":
            "开放问题与待验证决策",
        "TryRevive 的最强用户是非技术 Maker、ADHD/执行功能困难用户，还是学生创业者？三者不可同时作为首发。":
            "首轮定位决定：主样本只选非技术学生/学生创业者；独立创作者与执行功能困难用户后续单独验证，不混合统计。",
        "用户需要 AI 给第一步，还是更需要真人/同伴的承诺机制？":
            "待验证：用户主要需要 AI 给第一步，还是更需要真人/同伴承诺；在 concierge 中记录两种方式的启动率和完成率。",
        "“完成证据”会增加动力还是增加负担？":
            "待验证：完成证据会增加动力还是负担；默认提供轻量、可选证据，并比较必填与可选方案的完成率和放弃率。",
    }
    for old, new in replacements.items():
        replace_exact(doc, old, new)

    insert_many_after(
        paragraph_by_text(doc, "证据上传或链接"),
        [
            "建议依据与上下文来源提示",
            "短冲刺模式（首个微动作后，最多 4 个 45 分钟阶段；每阶段一个交付物）",
        ],
    )
    insert_after(
        paragraph_by_text(doc, "日历/提醒"),
        "可选真人/同伴承诺（先在 concierge 中验证，不默认开启）",
    )
    insert_after(
        paragraph_by_text(doc, "首个动作生成时间：从 intake 完成到看到动作的中位秒数。"),
        "首个真实动作启动时间：从看到动作到打开真实工具或开始计时的中位秒数。",
    )
    insert_many_after(
        paragraph_by_text(doc, "过大动作率：用户因“仍然太大”请求再次拆分的比例。"),
        [
            "证据负担率：用户因证据要求而跳过、放弃或请求移除的比例。",
            "计划停留率：创建 Brief 后继续规划但未在 30 分钟内开始真实动作的比例。",
        ],
    )
    insert_many_after(
        paragraph_by_text(doc, "名称混淆：TryRevive 可能与健康/影像产品混淆。缓解：公开发布前检查域名、App Store、商标和平台政策。"),
        [
            "首发人群过宽：混合样本会掩盖真实使用场景。缓解：首轮只统计同一主样本，其他人群单独记录。",
            "证据负担：完成证明可能把动力工具变成审查工具。缓解：证据默认轻量、可选，并监控证据负担率。",
            "动力机制误判：用户可能更需要真人/同伴承诺而不是 AI 建议。缓解：在 concierge 中做机制对照，验证后再产品化。",
            "渠道偏差：Hackathon 或公开构建的关注不等于留存。缓解：只把它们当招募和交付期限，决策只看外部用户行为。",
        ],
    )

    insert_improvement_section(doc)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
